"""把批次筛选结果投影为与页面一致的无模板 XLSX，不读取或修改数据库。"""

from __future__ import annotations

from datetime import timezone
from io import BytesIO
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = (
    "序号",
    "标题",
    "链接",
    "来源",
    "作者",
    "发布时间",
    "可见状态",
    "舆情结果",
    "评论数",
    "点赞数",
)
ANALYSIS_NAMES = {
    "analysis_queued": "等待分析",
    "analysis_running": "分析中",
    "analysis_completed": "分析成功",
    "analysis_partial": "分析不完整",
    "analysis_failed": "分析失败",
    "analysis_paused": "分析暂停",
    "analysis_disabled": "分析禁用",
}
SENTIMENT_NAMES = {"negative": "负面", "non_negative": "非负面", "unrelated": "不相关"}
SOURCE_NAMES = {"ai": "AI", "manual": "人工", "inherited_manual": "继承人工"}


def render_filtered_table(posts: list[dict[str, Any]]) -> bytes:
    """按传入的完整顺序写出列表字段；文本强制为字符串，缺失值不补零。"""
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("筛选结果")
    sheet.freeze_panes = "A2"
    for index, width in enumerate((8, 60, 55, 32, 24, 22, 14, 28, 12, 12), 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    header = []
    for label in HEADERS:
        cell = WriteOnlyCell(sheet, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="365C50")
        header.append(cell)
    sheet.append(header)
    for index, post in enumerate(posts, 1):
        source = post.get("source_name") or post.get("circle_name") or ""
        if post.get("list_order_name"):
            source = f"{source}（{post['list_order_name']}）"
        raw_status = post.get("raw_status")
        if isinstance(raw_status, dict) and raw_status.get("cross_forum_aggregate") is True:
            source += "（跨论坛）"
        sentiment = SENTIMENT_NAMES.get(post.get("sentiment_result"))
        if sentiment:
            origin = SOURCE_NAMES.get(post.get("sentiment_source"))
            if origin:
                sentiment = f"{sentiment}（{origin}）"
        else:
            sentiment = ANALYSIS_NAMES.get(post.get("analysis_status"))
        visibility = {"visible": "可见", "hidden": "不可见", "unknown": "未知"}.get(
            post.get("visibility")
        )
        if post.get("is_deleted"):
            visibility, sentiment = "删除", "已跳过（帖子已删除）"
        published = post.get("published_at")
        if published:
            if published.tzinfo is None:
                published = published.replace(tzinfo=timezone.utc)
            published = published.astimezone(ZoneInfo("Asia/Shanghai")).replace(tzinfo=None)
        values = (
            index,
            post.get("title"),
            post.get("url"),
            source,
            post.get("author"),
            published,
            visibility,
            sentiment,
            post.get("reply_count"),
            post.get("like_count"),
        )
        cells = []
        for column, value in enumerate(values, 1):
            cell = WriteOnlyCell(sheet)
            cell.value = ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value
            if isinstance(value, str):
                cell.data_type = "s"
            if column == 6:
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cells.append(cell)
        sheet.append(cells)
    sheet.auto_filter.ref = f"A1:J{len(posts) + 1}"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
