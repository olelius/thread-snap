"""新增模板字段、三平台来源、完整PNG及导出缓存的定向验证。"""

import hashlib
import json
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw
from sqlalchemy import select

from tests import test_backend
from tests.test_backend import AppCase
from threadsnap.errors import DomainError
from threadsnap.models import (
    Circle,
    CircleTask,
    ManualSentimentRevision,
    PostSnapshot,
    ScreenshotArtifactContribution,
    ScreenshotArtifactGroup,
    ScreenshotArtifactVersion,
    SentimentAnalysis,
)
from threadsnap.template_fields import EXTRA_FIELDS
from threadsnap.templates import FIELD_REGISTRY


class TemplateExtensionTests(AppCase):
    _seed_completed_run = test_backend.TemplateTests._seed_completed_run

    def template(self, circle, fields, *, conflict=False):
        """复用真实上传校验和模板版本，不跳过字段解析。"""
        workbook = Workbook()
        sheet = workbook.active
        for column, field in enumerate(fields, 1):
            sheet.cell(1, column, FIELD_REGISTRY[field]["description"])
            sheet.cell(2, column, f"s.{circle.export_key}.{field}")
        if conflict:
            sheet.cell(3, len(fields), "保留内容")
        path = Path(self.temp.name) / "template.xlsx"
        workbook.save(path)
        return self.container.templates.upload("扩展模板", path.name, path.read_bytes())

    def export(self, run, version):
        record = self.container.templates.create_export(run.id, version["version_id"])
        return record, self.container.templates.export_path(record["id"])

    def screenshot_group(self, run, circle):
        """建立两张独立完整框选PNG及对应成果版本，不拼接测试输入。"""
        tiles = []
        for index, size in enumerate(((800, 1000), (400, 300))):
            image = Image.new("RGB", size, "white")
            ImageDraw.Draw(image).rectangle((20, 30, size[0] - 20, 100), outline="red", width=5)
            path = Path(self.temp.name) / f"page-{index}.png"
            image.save(path)
            tiles.append({"index": index, "path": str(path), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "width": size[0], "height": size[1]})
        with self.container.sessions.begin() as db:
            task = db.scalar(select(CircleTask).where(CircleTask.run_id == run.id))
            group = ScreenshotArtifactGroup(chain_root_run_id=run.id, platform_code=task.platform_code, external_id=task.external_id, section=task.section, list_order=task.list_order, status="ready", dirty=False, current_version=1, item_count=2, negative_count=2)
            db.add(group)
            db.flush()
            db.add(ScreenshotArtifactContribution(group_id=group.id, run_id=run.id, circle_task_id=task.id))
            db.add(ScreenshotArtifactVersion(group_id=group.id, version=1, status="ready", reason="test", input_sha256="a" * 64, item_count=2, negative_count=2, tiles=tiles, items=[{"tile_index": i, "sentiment_result": "negative"} for i in range(2)], package_path="unused.zip", package_sha256="b" * 64))
            return group.id, tiles

    def test_preserves_old_registry_and_exposes_same_fields_for_three_platforms(self):
        legacy = {key: value for key, value in FIELD_REGISTRY.items() if key not in EXTRA_FIELDS}
        digest = hashlib.sha256(json.dumps(legacy, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
        self.assertEqual(digest, "e80bb37bd362589f087a9ccfef00c46fbb52a00eee050d3c20a99eb1b3dd1dc4")
        self.assertEqual(len(FIELD_REGISTRY), 51)
        with self.container.sessions.begin() as db:
            circles = [Circle(platform_code=code, external_id="same", name="同名来源", url=f"https://example.test/{code}", source_kind="configured") for code in ("dongchedi", "autohome", "yiche")]
            db.add_all(circles)
            db.flush()
        tags = []
        for circle in circles:
            response = self.client.get("/api/v1/template-fields", params={"source_id": circle.id})
            self.assertEqual(response.status_code, 200)
            fields = response.json()
            self.assertEqual(len(fields), 51)
            self.assertEqual({entry["field"] for entry in fields}, set(FIELD_REGISTRY))
            tags.extend(entry["tag"] for entry in fields)
        self.assertEqual(len(set(tags)), 153)

    def test_all_new_scalar_fields_manual_priority_and_deleted_post(self):
        run, circle = self._seed_completed_run()
        with self.container.sessions.begin() as db:
            posts = db.scalars(select(PostSnapshot).where(PostSnapshot.run_id == run.id).order_by(PostSnapshot.order_index)).all()
            post = posts[0]
            post.analysis_status, post.sentiment_result, post.sentiment_source = "analysis_completed", "negative", "manual"
            post.sentiment_updated_at = datetime(2026, 9, 13, 2, 30, tzinfo=timezone.utc)
            db.add(SentimentAnalysis(post_id=post.id, platform_code="dongchedi", platform_post_id=post.platform_post_id, input_hash="a" * 64, status="analysis_completed", config_revision=1, subject_version=1, model_code="test-model", result="non_negative", summary="=原有AI总结", matched_subjects=["车型甲"], primary_category="other", duration_ms=0, modalities={"text": {"evidence": ["文字事实"]}, "image": {"items": [{"input_index": 0, "evidence": ["图片事实"]}]}}))
            db.add(ManualSentimentRevision(post_id=post.id, action="set_result", result="negative", primary_category="product_complaint", secondary_categories=["service_complaint"], note="=人工说明"))
            posts[1].raw_status = {"content_state": "deleted"}
            posts[1].sentiment_result = "negative"
            posts[1].analysis_status = "analysis_disabled"
        fields = [key for key in EXTRA_FIELDS if key != "source.screenshot"]
        version = self.template(circle, fields)
        record, path = self.export(run, version)
        sheet = load_workbook(path).active
        cells = {field: sheet.cell(2, index + 1) for index, field in enumerate(fields)}
        self.assertEqual(cells["platform.name"].value, "懂车帝")
        self.assertEqual(cells["run.number"].value, run.number)
        self.assertEqual(cells["sentiment.result_name"].value, "负面")
        self.assertEqual(cells["sentiment.source_name"].value, "人工")
        self.assertEqual(cells["sentiment.primary_category"].value, "product_complaint")
        self.assertEqual(cells["sentiment.duration_ms"].value, 0)
        self.assertEqual(cells["sentiment.manual_note"].data_type, "s")
        self.assertEqual(cells["sentiment.summary"].data_type, "s")
        self.assertEqual(cells["sentiment.updated_at"].value.hour, 10)
        self.assertIn("图片事实", cells["sentiment.image_evidence"].value)
        self.assertIsNone(sheet.cell(3, fields.index("sentiment.result") + 1).value)
        self.assertIsNone(sheet.cell(3, fields.index("sentiment.summary") + 1).value)
        self.assertTrue(sheet.cell(3, fields.index("post.is_deleted") + 1).value)
        self.assertEqual(record["id"], self.export(run, version)[0]["id"])
        before = path.read_bytes()
        with self.container.sessions.begin() as db:
            post = db.get(PostSnapshot, post.id)
            post.sentiment_result, post.sentiment_source = "non_negative", "ai"
        updated, _ = self.export(run, version)
        self.assertNotEqual(record["id"], updated["id"])
        self.assertEqual(path.read_bytes(), before)

    def test_screenshot_embeds_independent_original_bytes_and_preserves_post_rows(self):
        run, circle = self._seed_completed_run()
        _group_id, tiles = self.screenshot_group(run, circle)
        version = self.template(circle, ["post.title", "source.screenshot"])
        record, path = self.export(run, version)
        sheet = load_workbook(path).active
        self.assertEqual(sheet["A2"].value, "标题3001")
        self.assertEqual(sheet["A3"].value, "标题3002")
        self.assertEqual(len(sheet._images), 2)
        self.assertEqual([image.anchor._from.row for image in sheet._images], [1, 2])
        self.assertEqual(sheet.column_dimensions["B"].width, 18)
        self.assertLess(sheet._images[0].anchor.ext.cx, 800 * 9525)
        self.assertEqual(sheet.row_dimensions[2].height, 120)
        for image in sheet._images:
            self.assertEqual(image.anchor._from.col, 1)
            self.assertLessEqual(sheet.row_dimensions[image.anchor._from.row + 1].height, 409)
        with zipfile.ZipFile(path) as bundle:
            media = [bundle.read(name) for name in bundle.namelist() if name.startswith("xl/media/")]
            self.assertCountEqual([hashlib.sha256(data).hexdigest() for data in media], [tile["sha256"] for tile in tiles])
        self.assertEqual(record["id"], self.export(run, version)[0]["id"])

    def test_screenshot_pending_missing_and_current_version_invalidate_cache(self):
        run, circle = self._seed_completed_run()
        group_id, tiles = self.screenshot_group(run, circle)
        version = self.template(circle, ["source.screenshot"])
        first, original = self.export(run, version)
        before = original.read_bytes()
        with self.container.sessions.begin() as db:
            db.get(ScreenshotArtifactGroup, group_id).dirty = True
        pending, path = self.export(run, version)
        self.assertNotEqual(first["id"], pending["id"])
        self.assertIn("尚未就绪", load_workbook(path).active["A2"].value)
        with self.container.sessions.begin() as db:
            db.get(ScreenshotArtifactGroup, group_id).dirty = False
        Path(tiles[0]["path"]).unlink()
        missing, path = self.export(run, version)
        self.assertNotEqual(first["id"], missing["id"])
        self.assertEqual(load_workbook(path).active["A2"].value, "截图文件缺失")
        self.assertEqual(len(load_workbook(path).active._images), 1)
        self.assertEqual(original.read_bytes(), before)

    def test_screenshot_range_conflict_preserves_template(self):
        run, circle = self._seed_completed_run()
        self.screenshot_group(run, circle)
        version = self.template(circle, ["post.title", "source.screenshot"], conflict=True)
        with self.assertRaises(DomainError) as error:
            self.export(run, version)
        self.assertEqual(error.exception.code, "EXPORT_RANGE_CONFLICT")
