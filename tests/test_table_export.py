"""筛选表格 XLSX 的真实接口、批次边界和字段一致性回归。"""

from __future__ import annotations

import hashlib
import json
from datetime import datetime, timezone
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import select

from tests.test_backend import AppCase
from threadsnap.db import Base
from threadsnap.models import Circle, CircleTask, ExtractionRun, PostSnapshot, Vehicle
from threadsnap.services import task_source_key
from threadsnap.table_export import HEADERS, render_filtered_table


def seed_export_fixture(container) -> dict:
    """建立两个选中来源、一个排除来源、关联补提和无关批次，供接口及浏览器共用。"""
    with container.sessions.begin() as db:
        runs = []
        for index in range(3):
            run = ExtractionRun(
                number=f"20260911-160000-00{index + 1}",
                trigger_type="manual",
                status="success",
                idempotency_scope="export-test",
                idempotency_key=str(index),
                request_hash="x" * 64,
                planned_count=66,
                completed_count=66,
                related_run_id=runs[0].id if index == 1 else None,
            )
            db.add(run)
            db.flush()
            runs.append(run)
        circles = []
        for index in range(3):
            vehicle = Vehicle(name=f"来源{index + 1}当前名称")
            db.add(vehicle)
            db.flush()
            circle = Circle(
                platform_code="dongchedi",
                external_id=str(40000 + index),
                name=f"圈子{index + 1}",
                url=f"https://example.test/circle/{index}",
                vehicle_id=vehicle.id,
                validation_status="verified",
            )
            db.add(circle)
            db.flush()
            circles.append(circle)
        tasks = []
        for index, (run, circle) in enumerate(
            [
                (runs[0], circles[0]),
                (runs[0], circles[1]),
                (runs[0], circles[2]),
                (runs[1], circles[0]),
                (runs[2], circles[0]),
            ]
        ):
            task = CircleTask(
                run_id=run.id,
                circle_id=circle.id,
                platform_code="dongchedi",
                external_id=circle.external_id,
                circle_name=circle.name,
                circle_url=circle.url,
                status="success",
                queue_sequence=index + 1,
                source_position=index if index < 3 else 0,
                list_order="latest_publish" if index == 1 else "latest_reply",
                target_count=60,
                completed_count=60,
                config_snapshot={"source_name": "创建时旧名称"},
            )
            db.add(task)
            db.flush()
            tasks.append(task)

        def add_post(task, index, **changes):
            """缺省为符合组合筛选的负面记录；按场景覆盖单个排除条件。"""
            values = dict(
                run_id=task.run_id,
                circle_task_id=task.id,
                platform_post_id=str(50000 + index),
                url=f"https://example.test/post/{index}",
                title=f"筛选样本{index:03d}",
                author="作者甲",
                published_at=datetime(2026, 9, 11, 2, 30, tzinfo=timezone.utc),
                content="正文不属于列表导出",
                visibility="visible",
                order_index=index,
                reply_count=index,
                like_count=100 - index,
                analysis_status="analysis_completed",
                sentiment_result="negative",
                sentiment_source="ai",
                raw_status={"diagnostic": "不应导出"},
            )
            values.update(changes)
            post = PostSnapshot(**values)
            db.add(post)
            db.flush()
            return post

        for index in range(60):
            task = tasks[3] if index == 59 else tasks[index % 2]
            add_post(task, index)
        add_post(tasks[3], 0, title="筛选样本重复补提")
        add_post(tasks[4], 900, title="筛选样本无关批次")
        add_post(tasks[2], 901, title="筛选样本未选来源")
        add_post(tasks[0], 902, title="其他标题")
        add_post(tasks[0], 903, visibility="hidden")
        add_post(tasks[0], 904, sentiment_result="non_negative")
        add_post(tasks[0], 905, analysis_status="analysis_partial")
        special = add_post(
            tasks[0],
            906,
            title='=HYPERLINK("https://example.test","测试")',
            author="+SUM(1,2)",
            like_count=None,
            reply_count=0,
            published_at=None,
            raw_status={"content_state": "deleted"},
        )
        return {
            "run_id": runs[0].id,
            "retry_id": runs[1].id,
            "other_id": runs[2].id,
            "special_id": special.id,
            "source_keys": [task_source_key(t) for t in tasks[:2]],
        }


class FilteredTableExportTests(AppCase):
    def setUp(self):
        super().setUp()
        self.fixture = seed_export_fixture(self.container)
        self.path = f"/api/v1/runs/{self.fixture['run_id']}/posts"
        self.filters = {
            "title": "筛选样本",
            "source_key": self.fixture["source_keys"],
            "visibility": "visible",
            "sentiment_result": "negative",
            "analysis_status": "analysis_completed",
            "sort_by": "reply_count",
            "sort_direction": "desc",
        }

    def snapshot(self):
        """比较所有业务表，确认导出未写入任何快照或导出记录。"""
        with self.container.sessions() as db:
            values = {
                table.name: [list(row) for row in db.execute(select(table))]
                for table in Base.metadata.sorted_tables
            }
        return hashlib.sha256(json.dumps(values, default=str, sort_keys=True).encode()).hexdigest()

    def test_all_filtered_pages_match_workbook_and_database_unchanged(self):
        before = self.snapshot()
        pages = [
            self.client.get(
                self.path, params={**self.filters, "offset": offset, "limit": 20}
            ).json()
            for offset in (0, 20, 40)
        ]
        expected = [post for page in pages for post in page["items"]]
        self.assertEqual([60, 60, 60], [page["total"] for page in pages])
        response = self.client.get(self.path + "/export", params=self.filters)
        self.assertEqual(
            200, response.status_code, response.text[:200] if response.status_code != 200 else ""
        )
        self.assertEqual("no-store", response.headers["cache-control"])
        self.assertIn("attachment; filename*=UTF-8''", response.headers["content-disposition"])
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.active
        rows = list(sheet.values)
        self.assertEqual(HEADERS, rows[0])
        self.assertEqual(61, len(rows))
        self.assertEqual([p["url"] for p in expected], [row[2] for row in rows[1:]])
        for index, (post, row) in enumerate(zip(expected, rows[1:], strict=True), 1):
            self.assertEqual(
                (
                    index,
                    post["title"],
                    post["url"],
                    f"{post['source_name']}（{post['list_order_name']}）",
                    post["author"],
                    datetime(2026, 9, 11, 10, 30),
                    "可见",
                    "负面（AI）",
                    post["reply_count"],
                    post["like_count"],
                ),
                row,
            )
        self.assertEqual("A2", sheet.freeze_panes)
        self.assertEqual("A1:J61", sheet.auto_filter.ref)
        self.assertEqual(before, self.snapshot())

    def test_each_filter_sort_and_retry_entry_reuses_list_contract(self):
        for filters in (
            {"title": "其他标题"},
            {"source_key": self.fixture["source_keys"]},
            {"visibility": "hidden"},
            {"sentiment_result": "non_negative"},
            {"analysis_status": "analysis_partial"},
            *(
                {"sort_by": field, "sort_direction": direction}
                for field in ("source", "published_at", "reply_count", "like_count")
                for direction in ("asc", "desc")
            ),
        ):
            with self.subTest(filters=filters):
                expected = self.client.get(self.path, params={**filters, "limit": 500}).json()[
                    "items"
                ]
                response = self.client.get(self.path + "/export", params=filters)
                self.assertEqual(200, response.status_code)
                rows = list(load_workbook(BytesIO(response.content)).active.values)[1:]
                self.assertEqual([p["url"] for p in expected], [row[2] for row in rows])
        response = self.client.get(
            f"/api/v1/runs/{self.fixture['retry_id']}/posts/export", params=self.filters
        )
        self.assertEqual(61, len(list(load_workbook(BytesIO(response.content)).active.values)))

    def test_zero_missing_deleted_and_formula_text(self):
        response = self.client.get(self.path + "/export", params={"title": "HYPERLINK"})
        sheet = load_workbook(BytesIO(response.content), data_only=False).active
        self.assertEqual("s", sheet["B2"].data_type)
        self.assertTrue(sheet["B2"].value.startswith("=HYPERLINK"))
        self.assertEqual("s", sheet["E2"].data_type)
        self.assertEqual("+SUM(1,2)", sheet["E2"].value)
        self.assertIsNone(sheet["F2"].value)
        self.assertEqual("删除", sheet["G2"].value)
        self.assertEqual("已跳过（帖子已删除）", sheet["H2"].value)
        self.assertEqual(0, sheet["I2"].value)
        self.assertIsNone(sheet["J2"].value)

    def test_status_empty_missing_and_invalid_parameters(self):
        self.assertEqual(404, self.client.get("/api/v1/runs/missing/posts/export").status_code)
        empty = self.client.get(self.path + "/export", params={"title": "无匹配结果"})
        self.assertEqual((409, "EXPORT_EMPTY"), (empty.status_code, empty.json()["code"]))
        for status in ("queued", "running", "waiting_for_auth", "partial_success", "failed"):
            with self.subTest(status=status):
                with self.container.sessions.begin() as db:
                    db.get(ExtractionRun, self.fixture["run_id"]).status = status
                response = self.client.get(self.path + "/export")
                self.assertEqual(
                    200 if status in ("partial_success", "failed") else 409, response.status_code
                )
        self.assertEqual(
            422, self.client.get(self.path + "/export", params={"sort_by": "invalid"}).status_code
        )

    def test_recurring_run_and_pagination_do_not_truncate_export(self):
        with self.container.sessions.begin() as db:
            db.get(ExtractionRun, self.fixture["run_id"]).trigger_type = "recurring"
        response = self.client.get(
            self.path + "/export", params={**self.filters, "limit": 20, "offset": 40}
        )
        self.assertEqual(61, len(list(load_workbook(BytesIO(response.content)).active.values)))

    def test_export_has_no_list_api_500_row_cap(self):
        with self.container.sessions.begin() as db:
            task = db.scalar(select(CircleTask).where(CircleTask.run_id == self.fixture["run_id"]))
            for index in range(600):
                db.add(
                    PostSnapshot(
                        run_id=task.run_id,
                        circle_task_id=task.id,
                        platform_post_id=f"large-{index}",
                        url=f"https://example.test/large/{index}",
                        title=f"大结果集{index}",
                        order_index=1000 + index,
                        visibility="visible",
                    )
                )
        response = self.client.get(self.path + "/export", params={"title": "大结果集"})
        self.assertEqual(200, response.status_code)
        self.assertEqual(601, len(list(load_workbook(BytesIO(response.content)).active.values)))

    def test_renderer_analysis_states_unicode_and_control_characters(self):
        content = render_filtered_table(
            [
                {
                    "title": "中文\x01换行\n标题",
                    "analysis_status": "analysis_paused",
                    "source_name": "来源",
                    "list_order_name": "最新发布",
                    "raw_status": {"cross_forum_aggregate": True},
                },
                {"sentiment_result": "negative", "sentiment_source": "inherited_manual"},
            ]
        )
        sheet = load_workbook(BytesIO(content)).active
        self.assertEqual("中文换行\n标题", sheet["B2"].value)
        self.assertEqual("分析暂停", sheet["H2"].value)
        self.assertEqual("来源（最新发布）（跨论坛）", sheet["D2"].value)
        self.assertEqual("负面（继承人工）", sheet["H3"].value)
