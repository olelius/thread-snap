"""独立圈内内容数的解析、同次冻结、平台列组和历史边界回归。"""

from __future__ import annotations

import copy
import hashlib
import json
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

from lxml import html
from openpyxl import load_workbook
from PIL import Image
from sqlalchemy import select

from tests import test_reputation as existing
from tests import test_reputation_dongchedi_optional_metrics as optional
from tests.reputation_circle_fixtures import circle_html, circle_result_fields
from threadsnap.collectors.dongchedi import DongchediCollector, normalize_circle_url
from threadsnap.collectors.dongchedi_count import (
    circle_count_visible_text,
    parse_circle_content_count,
)
from threadsnap.errors import DomainError
from threadsnap.models import ReputationMappingValidationAttempt, ReputationResult
from threadsnap.reputation_adapter import ReputationAdapterError, ReputationMappingTarget
from threadsnap.reputation_dongchedi import DongchediReputationAdapter
from threadsnap.reputation_registry import REPUTATION_PLATFORMS


class CircleCountParserTests(unittest.TestCase):
    """复现两份脱敏真实源形状，并穷尽显式字段边界。"""

    def setUp(self):
        self.target = ReputationMappingTarget(
            "fixture-count",
            "4615",
            "https://www.dongchedi.com/auto/series/4615",
            "坦克300",
            "hash",
        )
        self.url = "https://www.dongchedi.com/community/4615/dongtai-release"

    def parse(self, content, *, target=None, url=None):
        return DongchediReputationAdapter._parse_circle_content(
            target or self.target, url or self.url, content
        )

    def payload(self, *, count=0, visible="共0条内容", **overrides):
        props = {
            "series_id": "4615",
            "cheyouHead": {"series_id": 4615, "series_name": "坦克300"},
            "cheyouList": {"total_count": count},
            **overrides,
        }
        return (
            f'<html><head><meta charset=utf-8></head><body><div>{visible}</div><script id="__NEXT_DATA__">'
            + json.dumps({"props": {"pageProps": props}}, ensure_ascii=False)
            + "</script></body></html>"
        ).encode()

    def test_two_real_source_shapes_and_zero_without_cards(self):
        for series_id, name, count in (
            ("4615", "坦克300", 36701),
            ("24729", "风云A9", 1882),
            ("4615", "坦克300", 0),
        ):
            with self.subTest(series_id=series_id, count=count):
                target = replace(
                    self.target, platform_vehicle_id=series_id, platform_display_name=name
                )
                url = f"https://www.dongchedi.com/community/{series_id}/dongtai-release"
                raw, proof = self.parse(circle_html(series_id, name, count), target=target, url=url)
                self.assertEqual(str(count), raw)
                self.assertEqual(
                    (count, count, f"共{count}条内容"),
                    (proof["json_count"], proof["visible_count"], proof["visible_raw"]),
                )
                self.assertEqual(url, proof["source_url"])
                self.assertEqual(64, len(proof["response_sha256"]))

    def test_missing_single_source_fallback_is_truthfully_recorded(self):
        for count, visible, expected in (
            (None, "", None),
            (12, "", "12"),
            (None, "共12条内容", "12"),
        ):
            with self.subTest(count=count, visible=visible):
                raw, proof = self.parse(self.payload(count=count, visible=visible))
                self.assertEqual(expected, raw)
                self.assertEqual(count, proof["json_count"])
                self.assertEqual(12 if visible else None, proof["visible_count"])

    def test_illegal_counts_corrupt_json_and_source_conflicts_fail(self):
        cases = [
            self.payload(count=value)
            for value in (True, False, -1, 1.5, "12", [], {}, float("nan"))
        ]
        cases += [
            self.payload(visible=f"共{value}条内容")
            for value in ("-1", "1.5", "1,234", "1.2万", "NaN")
        ]
        cases += [
            self.payload(count=12, visible="共13条内容"),
            self.payload(visible="共0条内容 共1条内容"),
        ]
        cases += [self.payload(cheyouList=value) for value in ([], False, "bad")]
        cases += [
            b"<html><head><meta charset=utf-8></head><body>empty</body></html>",
            b'<script id="__NEXT_DATA__">{</script>',
            b'<script id="__NEXT_DATA__">null</script>',
        ]
        for content in cases:
            with self.subTest(content=content), self.assertRaises(ReputationAdapterError) as raised:
                self.parse(content)
            self.assertEqual("REPUTATION_CIRCLE_CONTENT_INVALID", raised.exception.code)

    def test_each_identity_and_landed_order_mismatch_fails(self):
        for props in (
            {"series_id": "24729"},
            {"cheyouHead": {"series_id": 24729, "series_name": "坦克300"}},
            {"cheyouHead": {"series_id": 4615, "series_name": "风云A9"}},
        ):
            with self.subTest(props=props), self.assertRaises(ReputationAdapterError) as raised:
                self.parse(self.payload(**props))
            self.assertEqual("REPUTATION_CIRCLE_IDENTITY_MISMATCH", raised.exception.code)
        for url in (
            self.url.replace("4615", "24729"),
            self.url + "/2",
            "https://www.dongchedi.com/community/4615",
            "https://example.com/community/4615/dongtai-release",
        ):
            with self.subTest(url=url), self.assertRaises(ReputationAdapterError):
                self.parse(self.payload(), url=url)

    def test_shared_helper_ignores_post_and_script_text_in_config_collector(self):
        body = self.payload().replace(
            b"</body>",
            '<section class="community-card"><p>共999条内容</p></section><script>共888条内容</script><div hidden>共777条内容</div></body>'.encode(),
        )
        document = html.fromstring(body.decode())
        self.assertEqual(
            (0, "共0条内容"), parse_circle_content_count(circle_count_visible_text(document))
        )
        collector = object.__new__(DongchediCollector)
        collector._get = Mock(return_value=Mock(status_code=200, content=body, url=self.url))
        collector._detect_auth = Mock()
        collector._normalize_card_rows = Mock(return_value=[])
        result = collector._fetch_circle_page(self.url, 1, expected_count=None)
        self.assertEqual(0, result["total_count"])
        collector._get.assert_called_once_with(self.url)
        self.assertEqual(("4615", self.url), normalize_circle_url(self.url))

    def test_one_bounded_get_status_and_auth_contract(self):
        adapter = DongchediReputationAdapter(None, timeout_seconds=90)
        session = Mock()
        adapter._http_session = Mock(return_value=session)
        session.get.return_value = Mock(status_code=200, content=self.payload(), url=self.url)
        raw, _, _ = adapter._visit_circle_content(self.target, timeout_seconds=3)
        self.assertEqual("0", raw)
        session.get.assert_called_once_with(self.url, timeout=3)
        for status in (403, 429, 503):
            session.get.reset_mock()
            session.get.return_value = Mock(status_code=status, content=b"error", url=self.url)
            with self.assertRaises(ReputationAdapterError) as raised:
                adapter._visit_circle_content(self.target)
            self.assertEqual(status in (429, 503), raised.exception.retryable)
            self.assertEqual(1, session.get.call_count)
        session.get.reset_mock()
        with self.assertRaises(ReputationAdapterError):
            adapter._visit_circle_content(self.target, timeout_seconds=0)
        session.get.assert_not_called()
        session.get.return_value = Mock(status_code=200, content=b"login-required", url=self.url)
        with self.assertRaises(ReputationAdapterError) as raised:
            adapter._visit_circle_content(self.target)
        self.assertEqual("AUTH_REQUIRED", raised.exception.code)


class CircleCountVisitTests(unittest.IsolatedAsyncioTestCase):
    """两条既有访问路径都携带新计数，不变更评分截图数量。"""

    def setUp(self):
        self.fixture = optional.DongchediOptionalMetricTests()
        self.fixture.setUp()

    async def test_browser_extra_count_preserves_three_measurements_and_one_png(self):
        fixture = self.fixture
        fixture.adapter.include_circle_content_count = True
        browser, page, _ = fixture._browser()
        fields = circle_result_fields(fixture.target, 0)
        fixture.adapter._visit_circle_content = Mock(
            return_value=(
                "0",
                fields["circle_content_count_url"],
                fields["circle_content_count_measurement"],
            )
        )
        with tempfile.TemporaryDirectory() as temporary:
            result = await fixture.adapter._visit(browser, fixture.target, Path(temporary))
            self.assertEqual("0", result.circle_content_count_raw)
            self.assertEqual(3, len(result.measurements))
            page.screenshot.assert_awaited_once()
            self.assertEqual(1, fixture.adapter._visit_circle_content.call_count)
            self.assertLessEqual(
                fixture.adapter._visit_circle_content.call_args.kwargs["timeout_seconds"], 90
            )

    def test_http_extra_count_and_disabled_path(self):
        fixture = self.fixture
        fields = circle_result_fields(fixture.target, 23)
        fixture.adapter._visit_circle_content = Mock(
            return_value=(
                "23",
                fields["circle_content_count_url"],
                fields["circle_content_count_measurement"],
            )
        )
        result = fixture._http('{"props":{"pageProps":{}}}')
        self.assertIsNone(result.circle_content_count_raw)
        fixture.adapter._visit_circle_content.assert_not_called()
        fixture.adapter.include_circle_content_count = True
        result = fixture._http('{"props":{"pageProps":{}}}')
        self.assertEqual("23", result.circle_content_count_raw)
        self.assertEqual("37%", result.negative_rate_raw)
        self.assertEqual(1, fixture.adapter._visit_circle_content.call_count)


class CircleCountLifecycleTests(unittest.TestCase):
    """隔离数据库复用正式测试设置，覆盖API/冻结/门禁/派生产物。"""

    def setUp(self):
        self.fixture = existing.OfficialReputationLifecycleTest()
        self.fixture.setUp()
        self.addCleanup(self.fixture.tearDown)
        self.service = self.fixture.service
        self.client = self.fixture.client

    def test_validation_acceptance_and_official_keep_six_fields_and_sources(self):
        mapping = self.service.get_scope()["vehicles"][0]["mappings"]["dongchedi"]
        self.assertEqual("500", mapping["latest_metrics"]["circle_content_count"])
        self.assertIs(True, existing.OfficialFakeAdapter.last_include_circle_content_count)
        acceptance = self.service.create_real_acceptance([mapping["validation_run_id"]])
        metric = acceptance["results"][0]["metrics"]["circle_content_count"]
        self.assertTrue(
            self.service._circle_collection_proven(metric, mapping["platform_vehicle_id"])
        )
        self.assertEqual(
            ("500", "no_baseline", "neutral"),
            (metric["raw"], metric["comparison_status"], metric["tone"]),
        )
        due = self.service.check_schedule(self.fixture._at("2030-01-02", "10:00"))
        official = self.service.execute_run(due["queued_run_ids"][0])
        self.assertIs(True, existing.OfficialFakeAdapter.last_include_circle_content_count)
        self.assertEqual(
            (27, 27), (official["completed_count"], official["complete_evidence_count"])
        )
        self.assertTrue(
            all(
                len(row["metrics"]) == 6 and row["metrics"]["circle_content_count"]["source_url"]
                for row in official["results"]
            )
        )
        caps = self.client.get("/api/v1/reputation/capabilities").json()["reputation_platforms"]
        self.assertEqual([6, 6, 5], [len(item["supported_metrics"]) for item in caps])

    def test_new_acceptance_rejects_missing_source_proof_but_history_is_immutable(self):
        mapping = self.service.get_scope()["vehicles"][0]["mappings"]["dongchedi"]
        with self.service.sessions() as db:
            attempt = db.get(ReputationMappingValidationAttempt, mapping["validation_attempt_id"])
            metrics, gates = copy.deepcopy(attempt.metrics), copy.deepcopy(attempt.gate_results)
        for missing in ("field", "source", "proof", "option", "id", "count"):
            new_metrics, new_gates = copy.deepcopy(metrics), copy.deepcopy(gates)
            metric = new_metrics["frozen_metrics"]["circle_content_count"]
            if missing == "field":
                new_metrics["frozen_metrics"].pop("circle_content_count")
            elif missing == "source":
                metric["source_url"] = None
            elif missing == "proof":
                metric["source_measurement"] = None
            elif missing == "option":
                new_gates["collection_options"]["include_circle_content_count"] = False
            elif missing == "id":
                metric["source_measurement"]["platform_vehicle_id"] = "999"
            else:
                metric["raw"] = "999"
            with self.service.sessions.begin() as db:
                attempt = db.get(
                    ReputationMappingValidationAttempt, mapping["validation_attempt_id"]
                )
                attempt.metrics, attempt.gate_results = new_metrics, new_gates
            with self.subTest(missing=missing), self.assertRaises(DomainError) as raised:
                self.service.create_real_acceptance([mapping["validation_run_id"]])
            self.assertEqual("REPUTATION_ACCEPTANCE_COLLECTION_INCOMPLETE", raised.exception.code)
        with self.service.sessions.begin() as db:
            attempt = db.get(ReputationMappingValidationAttempt, mapping["validation_attempt_id"])
            attempt.metrics, attempt.gate_results = metrics, gates
        old = self.service.create_real_acceptance([mapping["validation_run_id"]])
        files = {
            kind: self.client.get(url).content
            for kind, url in old["downloads"].items()
            if kind in {"txt", "xlsx"}
        }
        with self.service.sessions.begin() as db:
            for result in db.scalars(
                select(ReputationResult).where(ReputationResult.run_id == old["id"])
            ):
                result.metrics = {
                    key: value
                    for key, value in result.metrics.items()
                    if key != "circle_content_count"
                }
            attempt = db.get(ReputationMappingValidationAttempt, mapping["validation_attempt_id"])
            attempt.metrics, attempt.gate_results = {}, {}
        historical = self.service.get_run(old["id"])
        self.assertNotIn("circle_content_count", historical["results"][0]["metrics"])
        self.assertEqual(
            historical, self.service.create_real_acceptance([mapping["validation_run_id"]])
        )
        for kind, body in files.items():
            self.assertEqual(body, self.client.get(historical["downloads"][kind]).content)

    def test_neutral_delta_zero_missing_and_retired_key_never_used(self):
        target = ReputationMappingTarget(
            "x", "4615", "https://www.dongchedi.com/auto/series/4615", "坦克300", "hash"
        )
        sample = existing.OfficialFakeAdapter(concurrency=2)
        with tempfile.TemporaryDirectory() as temporary:
            page = sample.validate_sync([target], Path(temporary) / "sample")[0]
        for count, baseline, direction in ((0, 2, "down"), (3, 2, "up"), (2, 2, "same")):
            current = replace(page, **circle_result_fields(target, count))
            metrics = self.service._official_metrics(
                current,
                {
                    "metrics": {
                        "circle_content_count": {"raw": str(baseline), "value": str(baseline)}
                    }
                },
            )
            self.assertEqual(
                (str(count), direction, "neutral"),
                (
                    metrics["circle_content_count"]["raw"],
                    metrics["circle_content_count"]["direction"],
                    metrics["circle_content_count"]["tone"],
                ),
            )
        metrics = self.service._official_metrics(
            page, {"metrics": {"circle_content": {"value": "500"}}}
        )
        self.assertEqual("no_baseline", metrics["circle_content_count"]["comparison_status"])
        metrics = self.service._official_metrics(
            replace(page, **circle_result_fields(target, None)), None
        )
        self.assertEqual("not_available", metrics["circle_content_count"]["comparison_status"])
        for code in ("autohome", "yiche"):
            self.assertEqual(
                set(REPUTATION_PLATFORMS[code].metric_keys), set(self.service._official_metrics(page, None, code))
            )

    def test_xlsx_dynamic_columns_neutral_color_and_single_remark_anchor(self):
        source = self.fixture.root / "evidence.png"
        Image.new("RGB", (100, 40), "white").save(source)
        original_hash = hashlib.sha256(source.read_bytes()).hexdigest()
        for codes, count, note in (
            (["dongchedi"], 11, 10),
            (["autohome"], 11, 10),
            (["yiche"], 10, 9),
            (["dongchedi", "autohome", "yiche"], 22, 21),
            (["yiche", "dongchedi"], 16, 15),
        ):
            run = SimpleNamespace(
                id="fixture-export", platform_codes=codes, planned_date="2030-01-02"
            )
            results, evidence = [], {}
            for code in codes:
                metrics = {
                    key: {"raw": "0", "tone": "neutral"}
                    for key in REPUTATION_PLATFORMS[code].metric_keys
                }
                result = SimpleNamespace(
                    id=code,
                    vehicle_id="v",
                    role="focus",
                    series_name="车系",
                    vehicle_name="车型",
                    platform_code=code,
                    metrics=metrics,
                    error_message=None,
                    evidence_required=code != "yiche",
                )
                results.append(result)
                if code != "yiche":
                    evidence[code] = SimpleNamespace(
                        id=code, metric_region_path=str(source), metric_region_sha256=original_hash
                    )
            path = self.fixture.root / ("-".join(codes) + ".xlsx")
            self.service._create_xlsx(run, results, evidence, path)
            sheet = load_workbook(path).active
            self.assertEqual(count, sheet.max_column)
            self.assertEqual("备注", sheet.cell(1, count).value)
            self.assertIsNone(sheet.cell(2, count).value)
            if evidence:
                self.assertEqual(1, len(sheet._images))
                self.assertEqual(note, sheet._images[0].anchor._from.col)
            for index, header in enumerate(next(sheet.iter_rows(values_only=True)), start=1):
                if "圈内内容数" in str(header):
                    self.assertEqual("0", sheet.cell(2, index).value)
                    self.assertTrue(sheet.cell(2, index).fill.fgColor.rgb.endswith("F8FAFC"))
                    self.assertEqual(
                        16, sheet.column_dimensions[sheet.cell(1, index).column_letter].width
                    )
            self.assertEqual(original_hash, hashlib.sha256(source.read_bytes()).hexdigest())

    def test_xlsx_remarks_follow_frozen_evidence_requirement_not_current_platform_mode(self):
        """当前URL平台的旧必需截图仍报缺失；未要求截图不因无PNG而报缺失。"""

        run = SimpleNamespace(
            id="fixture-missing-evidence",
            platform_codes=["dongchedi", "yiche"],
            planned_date="2030-01-02",
        )
        for required, expected in (
            (False, "懂车帝：证据缺失"),
            (True, "懂车帝：证据缺失；易车：证据缺失"),
        ):
            results = [
                SimpleNamespace(
                    id=code,
                    vehicle_id="v",
                    role="focus",
                    series_name="车系",
                    vehicle_name="车型",
                    platform_code=code,
                    metrics={
                        key: {"raw": "0", "tone": "neutral"}
                        for key in REPUTATION_PLATFORMS[code].metric_keys
                    },
                    error_message=None,
                    evidence_required=(True if code == "dongchedi" else required),
                )
                for code in run.platform_codes
            ]
            path = self.fixture.root / f"frozen-requirement-{required}.xlsx"
            self.service._create_xlsx(run, results, {}, path)
            sheet = load_workbook(path).active
            self.assertEqual(expected, sheet.cell(2, sheet.max_column).value)
        missing_result_path = self.fixture.root / "missing-result.xlsx"
        self.service._create_xlsx(run, [results[1]], {}, missing_result_path)
        sheet = load_workbook(missing_result_path).active
        self.assertEqual(
            "懂车帝：未取得执行结果；易车：证据缺失", sheet.cell(2, sheet.max_column).value
        )

    def test_report_and_txt_only_describe_actual_circle_quantity_change(self):
        due = self.service.check_schedule(self.fixture._at("2030-01-02", "10:00"))
        first = self.service.execute_run(due["queued_run_ids"][0])
        existing.OfficialFakeAdapter.circle_count_overrides = {"official-01": 0}
        due = self.service.check_schedule(self.fixture._at("2030-01-03", "10:00"))
        second = self.service.execute_run(due["queued_run_ids"][0])
        report = self.service.generate_report(second["id"])
        self.assertIn("圈内内容数0，较昨日下降500", report["report_text"])
        self.assertEqual(
            report["report_text"].encode(), self.client.get(report["downloads"]["txt"]).content
        )
        self.assertEqual("neutral", second["results"][0]["metrics"]["circle_content_count"]["tone"])
        self.assertEqual(first, self.service.get_run(first["id"]))


if __name__ == "__main__":
    unittest.main()
